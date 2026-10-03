"""Edge cases for document-fidelity, numbering and the register — lock behavior
and guard the error paths."""

import io

from docx import Document

from util import burn_superadmin, make_owner, auth


def _doc_import(client, token, acc, **form):
    d = Document(); d.add_heading("Doc", level=1)
    buf = io.BytesIO(); d.save(buf); buf.seek(0)
    data = {"file": (buf, "x.docx")}
    data.update(form)
    return client.post(f"/api/accounts/{acc}/protocols/import", data=data,
                       content_type="multipart/form-data", headers=auth(token))


# ── render.docx / download-original error paths ──────────────────────────────

def test_render_docx_400_for_steps_protocol_without_body(client):
    burn_superadmin(client)
    token, acc, _ = make_owner(client, "ec1@corp.com", "Eco")
    p = client.post(f"/api/accounts/{acc}/protocols",
                    json={"title": "Plain steps"}, headers=auth(token)).get_json()["protocol"]
    res = client.post(f"/api/accounts/{acc}/protocols/{p['id']}/render.docx",
                      json={"variables": {}}, headers=auth(token))
    assert res.status_code == 400


def test_download_original_404_when_none_stored(client):
    burn_superadmin(client)
    token, acc, _ = make_owner(client, "ec2@corp.com", "Eco")
    # Pasted-text document import keeps no original file.
    p = client.post(f"/api/accounts/{acc}/protocols/import",
                    json={"title": "Pasted", "text": "# Title\nbody", "mode": "document"},
                    headers=auth(token)).get_json()["protocol"]
    assert p["doc_format"] == "document" and p["has_original"] is False
    res = client.get(f"/api/accounts/{acc}/protocols/{p['id']}/original", headers=auth(token))
    assert res.status_code == 404


def test_render_docx_from_blocks_without_original(client):
    """A pasted document (no original file) still renders a .docx from its blocks."""
    burn_superadmin(client)
    token, acc, _ = make_owner(client, "ec3@corp.com", "Eco")
    p = client.post(f"/api/accounts/{acc}/protocols/import",
                    json={"title": "Pasted", "text": "# {{product}} Spec\nContent", "mode": "document"},
                    headers=auth(token)).get_json()["protocol"]
    res = client.post(f"/api/accounts/{acc}/protocols/{p['id']}/render.docx",
                      json={"variables": {"product": "CAR-T"}}, headers=auth(token))
    assert res.status_code == 200
    doc = Document(io.BytesIO(res.data))
    text = "\n".join(par.text for par in doc.paragraphs)
    assert "CAR-T" in text and "{{" not in text


# ── numbering edge cases ─────────────────────────────────────────────────────

def test_template_import_gets_no_document_number(client):
    burn_superadmin(client)
    token, acc, _ = make_owner(client, "ec4@corp.com", "Eco")
    p = _doc_import(client, token, acc, as_template="true", category="SOP",
                    doc_category="QA").get_json()["protocol"]
    assert p["is_template"] is True
    assert p["document_number"] == ""        # templates never consume a number


def test_recategorizing_keeps_existing_number(client):
    burn_superadmin(client)
    token, acc, _ = make_owner(client, "ec5@corp.com", "Eco")
    p = _doc_import(client, token, acc, doc_category="EQ").get_json()["protocol"]
    assert p["document_number"] == "EQ-001"
    # Changing category does not renumber an already-numbered controlled document.
    res = client.put(f"/api/accounts/{acc}/protocols/{p['id']}",
                     json={"doc_category": "QA"}, headers=auth(token))
    updated = res.get_json()["protocol"]
    assert updated["doc_category"] == "QA"
    assert updated["document_number"] == "EQ-001"


def test_numbers_are_isolated_per_account(client):
    burn_superadmin(client)
    t1, a1, _ = make_owner(client, "ec6a@corp.com", "A1")
    t2, a2, _ = make_owner(client, "ec6b@corp.com", "A2")
    p1 = _doc_import(client, t1, a1, doc_category="TM").get_json()["protocol"]
    p2 = _doc_import(client, t2, a2, doc_category="TM").get_json()["protocol"]
    assert p1["document_number"] == "TM-001"
    assert p2["document_number"] == "TM-001"   # each account has its own sequence


# ── register status filtering ────────────────────────────────────────────────

def test_register_status_filters(client):
    burn_superadmin(client)
    token, acc, _ = make_owner(client, "ec7@corp.com", "Eco")
    _doc_import(client, token, acc, doc_category="EQ")   # a draft
    # Default (effective) is empty; "all" includes the draft.
    eff = client.get(f"/api/accounts/{acc}/protocols/register", headers=auth(token)).get_json()
    allr = client.get(f"/api/accounts/{acc}/protocols/register?status=all", headers=auth(token)).get_json()
    assert eff["procedures"] == []
    assert any(g["code"] == "EQ" for g in allr["procedures"])
