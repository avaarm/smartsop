"""Document-fidelity import: a real controlled document (batch record, CMC doc,
form) is kept 1-to-1 — tables stay tables, headings keep levels, checkboxes are
recognized — instead of being flattened into a mangled step list. The original
file is stored so a filled copy exports byte-faithfully."""

import io

from docx import Document

from util import burn_superadmin, make_owner, auth


def _batch_record_docx():
    d = Document()
    d.add_heading("Cell Processing Facility Batch Record", level=1)
    t = d.add_table(rows=3, cols=2)
    t.cell(0, 0).text = "Document Number"; t.cell(0, 1).text = "Title"
    t.cell(1, 0).text = "EQ-002"; t.cell(1, 1).text = "Operation of Biological Safety Cabinets"
    t.cell(2, 0).text = "Lot No.:"; t.cell(2, 1).text = "{{lot_number}}"
    d.add_paragraph("☐ Copy is an accurate reproduction of the current approved version.")
    d.add_paragraph("Product: {{product}}")
    buf = io.BytesIO(); d.save(buf)
    return buf.getvalue()


def _import_document(client, token, acc, name="BR-B098-01.docx", mode=None):
    data = {"file": (io.BytesIO(_batch_record_docx()), name)}
    if mode:
        data["mode"] = mode
    return client.post(f"/api/accounts/{acc}/protocols/import", data=data,
                       content_type="multipart/form-data", headers=auth(token))


def test_file_import_defaults_to_document(client):
    burn_superadmin(client)
    token, acc, _ = make_owner(client, "df1@corp.com", "Dco")
    res = _import_document(client, token, acc)
    assert res.status_code == 201
    body = res.get_json()
    assert body["mode"] == "document"
    p = body["protocol"]
    assert p["doc_format"] == "document"
    assert p["has_original"] is True
    assert p["original_filename"] == "BR-B098-01.docx"


def test_tables_and_checkboxes_preserved(client):
    burn_superadmin(client)
    token, acc, _ = make_owner(client, "df2@corp.com", "Dco")
    p = _import_document(client, token, acc).get_json()["protocol"]
    blocks = p["body"]
    types = [b["type"] for b in blocks]
    assert "heading" in types and "table" in types and "checkbox" in types

    # The References table stays a table (rows × cells), not pipe-mangled text.
    table = next(b for b in blocks if b["type"] == "table")
    assert table["rows"][0] == ["Document Number", "Title"]
    assert table["rows"][1] == ["EQ-002", "Operation of Biological Safety Cabinets"]

    chk = next(b for b in blocks if b["type"] == "checkbox")
    assert chk["checked"] is False
    assert "accurate reproduction" in chk["text"]

    # No step rows were created — it's a document, not a step list.
    assert p["step_count"] == 0


def test_template_variables_scanned_from_body(client):
    burn_superadmin(client)
    token, acc, _ = make_owner(client, "df3@corp.com", "Dco")
    pid = _import_document(client, token, acc).get_json()["protocol"]["id"]
    res = client.get(f"/api/accounts/{acc}/protocols/{pid}/template-variables", headers=auth(token))
    assert res.status_code == 200
    variables = res.get_json()["variables"]
    assert "lot_number" in variables and "product" in variables


def test_render_docx_fills_variables_byte_faithfully(client):
    burn_superadmin(client)
    token, acc, _ = make_owner(client, "df4@corp.com", "Dco")
    pid = _import_document(client, token, acc).get_json()["protocol"]["id"]
    res = client.post(f"/api/accounts/{acc}/protocols/{pid}/render.docx",
                      json={"variables": {"lot_number": "L-2026-001", "product": "CAR-T Cells"}},
                      headers=auth(token))
    assert res.status_code == 200
    assert res.headers["Content-Type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml")
    filled = Document(io.BytesIO(res.data))
    text = "\n".join(p.text for p in filled.paragraphs)
    text += "\n" + "\n".join(c.text for tb in filled.tables for r in tb.rows for c in r.cells)
    assert "L-2026-001" in text and "CAR-T Cells" in text
    assert "{{" not in text          # every placeholder substituted


def test_download_original_returns_uploaded_file(client):
    burn_superadmin(client)
    token, acc, _ = make_owner(client, "df5@corp.com", "Dco")
    pid = _import_document(client, token, acc).get_json()["protocol"]["id"]
    res = client.get(f"/api/accounts/{acc}/protocols/{pid}/original", headers=auth(token))
    assert res.status_code == 200
    # It's a valid docx (zip) round-tripping the original content.
    doc = Document(io.BytesIO(res.data))
    assert any("Batch Record" in p.text for p in doc.paragraphs)


def test_copy_with_variables_fills_document_blocks(client):
    burn_superadmin(client)
    token, acc, _ = make_owner(client, "df6@corp.com", "Dco")
    pid = _import_document(client, token, acc).get_json()["protocol"]["id"]
    res = client.post(f"/api/accounts/{acc}/protocols/{pid}/copy",
                      json={"title": "Batch 42", "variables": {"lot_number": "L-42", "product": "MSC"}},
                      headers=auth(token))
    assert res.status_code == 201
    clone = res.get_json()["protocol"]
    assert clone["doc_format"] == "document"
    cells = [c for b in clone["body"] if b["type"] == "table" for row in b["rows"] for c in row]
    assert "L-42" in cells
    assert not any("{{lot_number}}" in c for c in cells)


def test_document_body_is_editable(client):
    burn_superadmin(client)
    token, acc, _ = make_owner(client, "df7@corp.com", "Dco")
    pid = _import_document(client, token, acc).get_json()["protocol"]["id"]
    new_body = [{"type": "heading", "level": 1, "text": "Edited Title"},
                {"type": "paragraph", "text": "Adapted for the new project."}]
    res = client.put(f"/api/accounts/{acc}/protocols/{pid}", json={"body": new_body}, headers=auth(token))
    assert res.status_code == 200
    assert res.get_json()["protocol"]["body"][0]["text"] == "Edited Title"
