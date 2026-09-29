"""Import a protocol from pasted text or an uploaded file.

Two paths, mirroring protocols.io's "Pasting new data" + AI importer:

- Deterministic splitting (numbered / lines / markdown) — no LLM needed.
- AI structuring — when Ollama is reachable, ask the model to turn freeform
  prose into titled steps; falls back to the numbered splitter on any failure.
"""

import io
import re
import logging

logger = logging.getLogger(__name__)

MAX_STEPS = 500

_NUMBERED = re.compile(r"^\s*(?:step\s*)?(\d+)\s*[.):]\s+(.*)", re.IGNORECASE)
_MD_HEADING = re.compile(r"^\s*#{1,6}\s+(.*)")
_BULLET = re.compile(r"^\s*[-*•]\s+(.*)")


def extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from an uploaded .docx / .pdf / .txt / .md file."""
    name = (filename or "").lower()
    if name.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs)
    if name.endswith(".pdf"):
        try:
            from pypdf import PdfReader
        except ImportError:
            from PyPDF2 import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    # .txt, .md, or unknown — decode as text
    return data.decode("utf-8", errors="replace")


# A hierarchical section number like "3.2.S.1", "3.2.P", or "4.1" — segments of
# digits or single letters joined by "." / "-", at least two deep. Recognizes
# CTD/CMC-style headings that carry no Word heading style. A plain "1. foo" (one
# segment) is NOT matched, so ordinary numbered steps aren't mistaken for headings.
_NUM_HEADING = re.compile(r"^\s*((?:\d+|[A-Za-z])(?:[.\-](?:\d+|[A-Za-z]))+)\.?\s+\S")


def _heading_level(style_name: str, text: str) -> int:
    """Return a heading level (1..6) for a paragraph, or 0 if it's body text.

    Word heading styles win; otherwise a hierarchical section number
    (e.g. "3.2.S.1 Nomenclature") is treated as a heading so structured
    documents keep their outline.
    """
    style = (style_name or "").lower()
    m = re.match(r"heading\s*(\d+)", style)
    if m:
        return min(int(m.group(1)), 6)
    if style in ("title",):
        return 1
    if style in ("subtitle",):
        return 2
    if len(text) <= 160:
        nm = _NUM_HEADING.match(text)
        if nm:
            return min(nm.group(1).replace("-", ".").count(".") + 1, 6)
    return 0


def _iter_docx_blocks(doc):
    """Yield paragraphs and tables in document order (so structure is preserved)."""
    from docx.document import Document as _DocT
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    parent = doc.element.body if isinstance(doc, _DocT) else doc
    for child in parent.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, doc)
        elif isinstance(child, CT_Tbl):
            yield Table(child, doc)


def _table_to_text(table) -> str:
    """Render a Word table as plain text, preserving rows/columns."""
    rows = []
    for row in table.rows:
        cells = [c.text.strip().replace("\n", " ") for c in row.cells]
        rows.append(" | ".join(cells))
    return "\n".join(r for r in rows if r.strip(" |"))


def extract_structured(filename: str, data: bytes):
    """Extract an uploaded doc into [{section, title, description}] preserving its
    structure 1-to-1: headings become sections/step titles, body paragraphs and
    tables are kept in full under their heading, in original order.

    Non-.docx files fall back to text + heading-aware line parsing.
    """
    name = (filename or "").lower()
    if name.endswith(".docx"):
        from docx import Document
        return _clean(_structured_from_docx(Document(io.BytesIO(data))))
    return _clean(_structured_from_text(extract_text(filename, data)))


def _structured_from_docx(doc):
    steps, current, h1 = [], None, ""
    from docx.text.paragraph import Paragraph as _Para

    def append_body(text):
        nonlocal current
        if current is None:
            current = {"section": h1, "title": "", "description": ""}
            steps.append(current)
        current["description"] = (current["description"] + ("\n" if current["description"] else "") + text)

    for block in _iter_docx_blocks(doc):
        if isinstance(block, _Para):
            text = block.text.strip()
            if not text:
                if current and current["description"] and not current["description"].endswith("\n"):
                    current["description"] += "\n"
                continue
            level = _heading_level(block.style.name if block.style else "", text)
            if level == 1:
                h1 = text
                current = {"section": text, "title": "", "description": ""}
                steps.append(current)
            elif level >= 2:
                current = {"section": h1, "title": text, "description": ""}
                steps.append(current)
            else:
                append_body(text)
        else:  # table
            tbl = _table_to_text(block)
            if tbl:
                append_body(tbl)
    return steps


def _structured_from_text(text: str):
    """Heading-aware structuring for pasted / non-docx text: markdown #/## and
    numbered section headings define the outline; everything else is body."""
    steps, current, h1 = [], None, ""
    for ln in (text or "").splitlines():
        stripped = ln.strip()
        md = _MD_HEADING.match(ln)
        if md:
            text_h = md.group(1).strip()
            hashes = len(ln) - len(ln.lstrip("#"))   # number of leading '#'
            level = 1 if hashes <= 1 else 2
        else:
            level = _heading_level("", stripped) if stripped else 0
            text_h = stripped
        if stripped and level == 1:
            h1 = text_h
            current = {"section": text_h, "title": "", "description": ""}
            steps.append(current)
        elif stripped and level >= 2:
            current = {"section": h1, "title": text_h, "description": ""}
            steps.append(current)
        elif stripped:
            if current is None:
                current = {"section": "", "title": "", "description": ""}
                steps.append(current)
            current["description"] = (current["description"] + ("\n" if current["description"] else "") + stripped)
    return steps


def _clean(steps):
    out = []
    for s in steps:
        section = (s.get("section") or "").strip()[:200]
        title = (s.get("title") or "").strip()[:500]
        desc = (s.get("description") or "").strip()
        if section or title or desc:
            out.append({
                "section": section,
                "title": title,
                "description": desc,
                "duration_seconds": s.get("duration_seconds"),
                "warning": (s.get("warning") or "").strip(),
            })
    return out[:MAX_STEPS]


# ── Document-fidelity extraction ──────────────────────────────────────────
# A controlled document (batch record, CMC section, form) is NOT a step list.
# It's headings, paragraphs, tables, approval blocks and fill-in fields. These
# helpers keep that structure 1-to-1 as ordered "blocks" so the document renders
# like the original instead of being flattened into mangled step titles.

_CHECK_UNCHECKED = "☐□❑⬜"          # empty checkbox glyphs
_CHECK_CHECKED = "☑☒✓✔■"           # ticked / filled checkbox glyphs


def _cell_text(cell) -> str:
    """All text in a table cell, paragraphs joined by newlines."""
    return "\n".join(p.text for p in cell.paragraphs).strip()


def _table_block(table) -> dict:
    """A Word table preserved as a grid of cell strings (rows × columns)."""
    rows = []
    for row in table.rows:
        rows.append([_cell_text(c) for c in row.cells])
    return {"type": "table", "rows": rows}


def _blocks_text(blocks) -> str:
    """Flatten blocks to plain text (for search / a description fallback)."""
    parts = []
    for b in blocks:
        if b.get("type") == "table":
            for row in b.get("rows", []):
                parts.append(" | ".join(row))
        else:
            parts.append(b.get("text", ""))
    return "\n".join(p for p in parts if p)


def extract_document(filename: str, data: bytes) -> dict:
    """Extract an uploaded file into faithful, structure-preserving blocks.

    Returns {"blocks": [...], "text": "..."} where each block is one of:
      {"type": "heading", "level": 1-6, "text": str}
      {"type": "paragraph", "text": str}
      {"type": "checkbox", "checked": bool, "text": str}
      {"type": "table", "rows": [[cell, ...], ...]}

    Tables stay tables, headings keep their level, checkboxes are recognized —
    so a batch record looks like the batch record, not a step list.
    """
    name = (filename or "").lower()
    if name.endswith(".docx"):
        from docx import Document
        return _document_from_docx(Document(io.BytesIO(data)))
    return _document_from_text(extract_text(filename, data))


def _document_from_docx(doc) -> dict:
    from docx.text.paragraph import Paragraph as _Para
    blocks = []
    for block in _iter_docx_blocks(doc):
        if isinstance(block, _Para):
            text = block.text.strip()
            if not text:
                continue
            first = text[0]
            if first in _CHECK_UNCHECKED or first in _CHECK_CHECKED:
                blocks.append({
                    "type": "checkbox",
                    "checked": first in _CHECK_CHECKED,
                    "text": text[1:].strip(),
                })
                continue
            level = _heading_level(block.style.name if block.style else "", text)
            if level:
                blocks.append({"type": "heading", "level": level, "text": text})
            else:
                blocks.append({"type": "paragraph", "text": text})
        else:  # table
            tb = _table_block(block)
            if any(any(c.strip() for c in row) for row in tb["rows"]):
                blocks.append(tb)
    return {"blocks": blocks, "text": _blocks_text(blocks)}


def _document_from_text(text: str) -> dict:
    """Blocks from pasted / non-docx text: headings + paragraphs (no tables)."""
    blocks = []
    for ln in (text or "").splitlines():
        s = ln.strip()
        if not s:
            continue
        md = _MD_HEADING.match(ln)
        if md:
            hashes = len(ln) - len(ln.lstrip("#"))
            blocks.append({"type": "heading", "level": min(hashes, 6), "text": md.group(1).strip()})
            continue
        level = _heading_level("", s)
        if level:
            blocks.append({"type": "heading", "level": level, "text": s})
        else:
            blocks.append({"type": "paragraph", "text": s})
    return {"blocks": blocks, "text": text or ""}


def split_into_steps(text: str, mode: str = "numbered"):
    """Split text into [{title, description}] using a deterministic strategy."""
    lines = (text or "").splitlines()

    if mode == "lines":
        return _clean([{"description": ln.strip()} for ln in lines if ln.strip()])

    if mode == "markdown":
        steps, current = [], None
        for ln in lines:
            heading = _MD_HEADING.match(ln)
            numbered = _NUMBERED.match(ln)
            bullet = _BULLET.match(ln)
            if heading:
                current = {"title": heading.group(1).strip(), "description": ""}
                steps.append(current)
            elif numbered:
                current = {"title": "", "description": numbered.group(2).strip()}
                steps.append(current)
            elif bullet:
                current = {"title": "", "description": bullet.group(1).strip()}
                steps.append(current)
            elif ln.strip() and current is not None:
                current["description"] = (current["description"] + "\n" + ln.strip()).strip()
            elif ln.strip():
                current = {"title": "", "description": ln.strip()}
                steps.append(current)
        return _clean(steps)

    # default: numbered — "1. foo", "2) bar", "Step 3: baz"; continuation lines append.
    steps, current, matched = [], None, False
    for ln in lines:
        m = _NUMBERED.match(ln)
        if m:
            matched = True
            current = {"title": "", "description": m.group(2).strip()}
            steps.append(current)
        elif ln.strip() and current is not None:
            current["description"] = (current["description"] + " " + ln.strip()).strip()
    if not matched:
        # No numbered list found — fall back to one step per line.
        return split_into_steps(text, "lines")
    return _clean(steps)


AI_SYSTEM = (
    "You convert a lab procedure into structured steps. Return ONLY a JSON object "
    '{"steps": [{"title": str, "description": str, "duration_minutes": number|null, '
    '"warning": str}]}. Keep each step atomic and preserve the original order.'
)


def ai_structure(text: str, ollama):
    """Use the LLM to structure text into steps. Returns list or None on failure."""
    try:
        if ollama is None or not ollama.check_health():
            return None
        result = ollama.generate_json(
            prompt=f"Convert this procedure into steps:\n\n{text[:8000]}",
            system_prompt=AI_SYSTEM,
        )
        raw = result.get("steps") if isinstance(result, dict) else result
        if not isinstance(raw, list):
            return None
        steps = []
        for s in raw:
            if not isinstance(s, dict):
                continue
            mins = s.get("duration_minutes")
            steps.append({
                "title": s.get("title", ""),
                "description": s.get("description", ""),
                "duration_seconds": int(mins * 60) if isinstance(mins, (int, float)) and mins else None,
                "warning": s.get("warning", ""),
            })
        return _clean(steps) or None
    except Exception as e:  # noqa: BLE001 — degrade gracefully to deterministic split
        logger.warning("AI structuring failed, falling back to split: %s", e)
        return None
